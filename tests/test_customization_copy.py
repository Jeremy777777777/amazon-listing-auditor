import json
from dataclasses import replace

from openpyxl import load_workbook

from listing_auditor.cli import _parser
from listing_auditor.compliance import check_compliance
from listing_auditor.compare import compare
from listing_auditor.copy_style import style_correction_findings
from listing_auditor.customization import option_groups, options
from listing_auditor.evidence import collect_evidence
from listing_auditor.models import AuditResult, Evidence, Finding, ListingRecord
from listing_auditor.report import write_reports


def record():
    return ListingRecord('Dell Pro 24 All-in-One QC24251 / 23.8" / Black / Intel i5-14500T / Intel Integrated Graphic Card / Bluetooth / 1920x1080 / WiFi / 60 Hz','VA-215','MegaPC','https://www.amazon.com/dp/B0H9V9BKXS','B0H9V9BKXS',1)


def capture():
    return {'asin':'B0H9V9BKXS','available':True,'source':'Observed Customize dialog','captured_at':'2026-10-19T20:29:13Z','options':{'ram':['16GB','32GB','64GB'],'storage':['0.5TB','1TB','2TB','4TB']},'base':{},'selected':{}}


def make_evidence():
    return Evidence(record().url,title='Old title: 8GB/16GB/32GB/64GB RAM, 256GB/512GB/1TB/2TB SSD',details={'RAM Memory Technology':'DDR5','Memory / RAM Memory Installed':'8 GB'},customization=capture())


def corrections(evidence):
    return style_correction_findings(record(),evidence,None,[Finding('model','QC24251','Old title','HIGH','Wrong title')])


def test_copy_uses_observed_options_not_old_title_or_base_attribute(tmp_path):
    evidence = make_evidence()
    findings = corrections(evidence)
    copy = {f.field:f.corrected_value for f in findings if f.field.startswith('copy_') and f.corrected_value}
    assert len(copy) == 3
    all_copy = '\n'.join(copy.values())
    for value in ('16GB','32GB','64GB','0.5TB','1TB','2TB','4TB'): assert value in all_copy
    for value in ('[[VERIFY','8GB/','256GB','512GB','1.3 GHz','FHD IR'): assert value not in all_copy
    assert len(copy['copy_bullet_points'].split('\n\n')) == 5
    assert len(copy['copy_title']) <= 200
    assert copy['copy_product_description'].startswith('<p><strong>')
    current = replace(evidence,title=copy['copy_title'],bullets=copy['copy_bullet_points'].split('\n\n'),description=copy['copy_product_description'])
    assert not any(f.rule_id == 'CPC-BASE-001' for f in check_compliance(record(),current))
    wrong = replace(current,title=copy['copy_title'].replace('0.5TB','512GB'))
    assert any(f.rule_id == 'CPC-BASE-001' and f.severity == 'HIGH' for f in check_compliance(record(),wrong))
    path = write_reports([AuditResult(record(),'FAIL',findings,evidence)],tmp_path)
    wb = load_workbook(path)
    assert wb.sheetnames[0] == '可复制文案'
    ready = wb['可复制文案']
    assert ready['B2'].value == copy['copy_title']
    assert ready['B3'].value == copy['copy_bullet_points']
    assert ready['B4'].value == copy['copy_product_description']
    assert wb['Product Details']['M6'].value == ready['B3'].value
    assert wb['Product Details']['L6'].value == ready['B4'].value
    assert any(c.value and '0.5TB' in str(c.value) for c in wb['Source Evidence']['E'])


def test_missing_options_omit_capacity_text_and_record_source_gap():
    evidence = replace(make_evidence(),customization={})
    findings = corrections(evidence)
    assert any(f.rule_id=='COPY-SOURCE-002' for f in findings)
    text = '\n'.join(f.corrected_value for f in findings if f.field.startswith('copy_'))
    assert '[[VERIFY' not in text
    assert '512GB' not in text and '64GB' not in text
    # A platform maximum or old title is not evidence that this is an offered option.
    assert '16GB' not in text


def test_invalid_units_wrong_asin_and_ambiguous_flags_are_rejected():
    evidence = make_evidence()
    for bad in ({**capture(),'asin':'B012345678'}, {**capture(),'available':'false'}, {**capture(),'options':{'ram':['16','32','64']}}, {**capture(),'source':''}):
        assert options(replace(evidence,customization=bad),'ram') == []
    assert options(evidence,'storage')[0] == '0.5TB'
    assert option_groups('8/16/32GB RAM')[0][0] == {(8.0,'GB'),(16.0,'GB'),(32.0,'GB')}
    assert option_groups('Choose 16GB, 32GB or 64GB DDR5 RAM.')[0][1] == 'ram'


def test_customization_evidence_can_be_loaded_by_cli_and_fixture(tmp_path):
    r=record()
    (tmp_path/f'{r.asin}.json').write_text(json.dumps({'url':r.url,'title':'Captured page'}),encoding='utf-8')
    custom=tmp_path/'choices.json'
    custom.write_text(json.dumps({'listings':{r.asin:capture()}}),encoding='utf-8')
    e=collect_evidence(r,tmp_path,custom)
    assert options(e,'ram')==['16GB','32GB','64GB']
    args=_parser().parse_args(['audit','--amazon-url',r.url,'--customization-evidence',str(custom)])
    assert args.customization_evidence==custom
    custom.write_text(json.dumps({**capture(),'asin':'B012345678'}),encoding='utf-8')
    assert options(collect_evidence(r,tmp_path,custom),'ram') == []


def test_original_base_capacity_is_not_compared_to_first_upgrade_choice():
    r=replace(record(),product_name=record().product_name+' / 8GB DDR5 RAM / 256GB SSD')
    e=replace(make_evidence(),title='MegaPC Custom All-in-One, Created Using Dell Pro 24 All-in-One QC24251, Intel Core i5-14500T, 16GB/32GB/64GB DDR5 RAM, 0.5TB/1TB/2TB/4TB SSD, Win 11 Pro')
    result=compare(r,e)
    assert not any(f.field in {'memory','storage','cpu'} and f.severity=='HIGH' for f in result.findings)
