from pathlib import Path

from openpyxl import load_workbook

from listing_auditor.compare import compare
from listing_auditor.erp import load_erp_record
from listing_auditor.evidence import collect_evidence
from listing_auditor.models import AuditResult, Evidence, ListingRecord
from listing_auditor.report import write_product_packages, write_reports


ROOT = Path(__file__).parents[1]


def _record() -> ListingRecord:
    return ListingRecord(
        product_name='Dell 15 DC15250 / Laptop / 15.6" / Touchscreen / Black / Intel Core i7-1355U / Intel UHD Graphics / Fingerprint / Bluetooth / 1920x1080 / Wifi / 60 Hz',
        sku="VL-1249",
        seller="MegaPC",
        url="https://www.amazon.com/dp/B0HBDTJNJV",
        asin="B0HBDTJNJV",
        row_number=34,
    )


def test_erp_export_confirms_internal_product_and_compliance_review() -> None:
    record = _record()
    erp = load_erp_record(record.sku, export_path=ROOT / "fixtures" / "erp" / "products.json")
    assert erp.available
    assert erp.fields["brand"] == "Dell"
    result = compare(record, collect_evidence(record, ROOT / "fixtures" / "listings"), erp)
    assert result.status == "FAIL"
    assert {finding.field for finding in result.findings} >= {
        "brand", "model", "cpu", "resolution", "display_size", "refresh_rate", "compliance_bullet_1_warranty"
    }


def test_excel_contains_only_findings(tmp_path: Path) -> None:
    record = _record()
    pass_result = AuditResult(record, "PASS", [], Evidence(record.url, title=record.product_name, source="test"))
    fail_result = compare(
        record,
        collect_evidence(record, ROOT / "fixtures" / "listings"),
        load_erp_record(record.sku, export_path=ROOT / "fixtures" / "erp" / "products.json"),
    )
    path = write_reports([pass_result, fail_result], tmp_path)
    workbook = load_workbook(path)
    sheet = workbook["问题清单"]
    assert sheet["A8"].value == "严重程度"
    assert sheet.max_row == 8 + len(fail_result.findings)
    assert all(sheet.cell(row=row, column=5).value == "VL-1249" for row in range(9, sheet.max_row + 1))


def test_product_package_uses_internal_id_folder(tmp_path: Path) -> None:
    record = _record()
    result = compare(
        record,
        collect_evidence(record, ROOT / "fixtures" / "listings"),
        load_erp_record(record.sku, export_path=ROOT / "fixtures" / "erp" / "products.json"),
    )
    paths = write_product_packages([result], tmp_path)
    package = tmp_path / "VL-1249"
    assert paths == [package / "listing-audit-review.xlsx"]
    assert {path.name for path in package.iterdir()} == {
        "listing-audit-review.xlsx", "report.json", "report.csv", "report.md", "source-evidence.json"
    }
