import hashlib
import json
from dataclasses import replace

from PIL import Image

from listing_auditor.compare import compare
from listing_auditor.compliance import check_compliance
from listing_auditor.copy_style import style_correction_findings
from listing_auditor.erp import load_erp_record
from listing_auditor.image_audit import audit_images
from listing_auditor.models import ERPRecord, Evidence, ListingRecord
from listing_auditor.policy import warranty


RECORD = ListingRecord("Dell Inspiron 15 / Laptop", "VL-1234", "MegaPC", "https://www.amazon.com/dp/B012345678", "B012345678", 1)


def clean_evidence():
    return Evidence(RECORD.url,
        title="MegaPC Custom Laptop, Created Using Dell Inspiron 15, Intel Core i7-1355U, 16GB RAM, 512GB SSD, Win 11 Pro",
        bullets=[warranty("Dell"), "Reliable Work — Intel Core i7-1355U.", "Display — 15.6 inch.", "Memory — 16GB RAM and 512GB SSD.", "System — Windows 11 Pro."],
        description="<strong>Warranty and Disclosure</strong><br/>" + warranty("Dell"),
        details={"Warranty Description": warranty("Dell")})


def test_fixed_catalog_and_manual_publication_are_separate():
    findings = check_compliance(RECORD, clean_evidence())
    assert {f.rule_id for f in findings} == {"CAT-PUBLICATION-001"}


def test_fixed_pro_does_not_allow_selectable_software_or_gpu():
    for text in ("Choose Windows 11 Pro software.", "Upgraded GPU for games.", "Windows customized by MegaPC."):
        assert any(f.rule_id == "CPC-SCOPE-001" for f in check_compliance(RECORD, replace(clean_evidence(), description=text)))
    assert not any(f.rule_id == "CPC-SCOPE-001" for f in check_compliance(RECORD, replace(clean_evidence(), description="Customized RAM and SSD. Fixed Windows 11 Pro upgrade. GPU remains factory configured.")))


def test_jtd_does_not_inherit_megapc_catalog_assertions():
    findings = check_compliance(replace(RECORD, seller="JTD"), replace(clean_evidence(), title="J-Tech Digital Customized Laptop, Created Using Dell Inspiron 15, 16GB RAM, 512GB SSD"))
    assert not any(f.rule_id.startswith("CAT-") for f in findings)


def test_correct_title_cannot_hide_wrong_description_specs():
    result = compare(RECORD, replace(clean_evidence(), description="Intel Core i5-1235U, 32GB RAM, 1TB SSD."))
    assert {f.field for f in result.findings if f.rule_id == "SPEC-CROSS-FIELD-001"} == {"product_description_cpu", "product_description_memory", "product_description_storage"}


def test_factory_os_is_not_sold_os_and_ssd_is_not_memory():
    erp = ERPRecord(RECORD.sku, {"os": "Windows 11 Home", "osScope": "factory", "ram": "16GB", "ssd": "512GB SSD"}, available=True)
    result = compare(RECORD, clean_evidence(), erp)
    assert not [f for f in result.findings if f.field in {"os", "memory", "storage"}]


def test_ambiguous_erp_export_requires_review(tmp_path):
    path = tmp_path / "erp.json"
    path.write_text(json.dumps([{"sku": RECORD.sku, "ram": "16GB"}, {"sku": RECORD.sku, "ram": "32GB"}]))
    record = load_erp_record(RECORD.sku, export_path=path)
    assert not record.available
    assert "Multiple" in record.error


def test_non_square_and_authentic_ocr_are_not_errors_but_reviews_need_hashes(tmp_path):
    image = tmp_path / "main.png"
    Image.new("RGB", (1200, 1000), "white").save(image)
    observations = tmp_path / "observations.json"
    item = {"slot": "MAIN", "sha256": hashlib.sha256(image.read_bytes()).hexdigest(), "ocr_text": "Dell", "licensed_provenance": "Seller photography"}
    observations.write_text(json.dumps({"images": [item]}))
    evidence = replace(clean_evidence(), image_urls=[str(image), str(image)])
    findings = audit_images(RECORD, evidence, observations_path=observations)
    assert not any(f.rule_id in {"IMG-MAIN-003", "IMG-TECH-004", "IMG-EVIDENCE-003"} for f in findings)
    assert any(f.field == "image_PT01_review_missing" for f in findings)
    item["sha256"] = "outdated"
    observations.write_text(json.dumps({"images": [item]}))
    assert any(f.rule_id == "IMG-EVIDENCE-003" for f in audit_images(RECORD, evidence, observations_path=observations))


def test_copy_does_not_invent_capacity_choices_and_escapes_html():
    evidence = replace(clean_evidence(), title="Wrong title")
    erp = ERPRecord(RECORD.sku, {"model": "Dell Inspiron 15", "category": "Laptop", "ram": "Up to 64GB RAM", "ssd": "1TB SSD", "cpu": "Intel Core i7-1355U", "screenSize": "15.6 inch <display>"}, available=True)
    findings = style_correction_findings(RECORD, evidence, erp, check_compliance(RECORD, evidence))
    copy = {f.field: f.corrected_value for f in findings}
    assert "[[VERIFY INSTALLED BASE RAM]]" in copy["copy_title"]
    assert "&lt;display&gt;" in copy["copy_product_description"]
    assert "choose" not in copy["copy_bullet_points"].casefold()


def test_semantic_overlap_uses_primary_fact_ids(tmp_path):
    path = tmp_path / "image.png"
    Image.new("RGB", (1000, 1100), "white").save(path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    observations = tmp_path / "observations.json"
    items = [{"slot": slot, "sha256": digest, "visual_review_100pct": True, "visual_review_200px": True, "primary_fact_ids": facts, "licensed_provenance": "Own photo"} for slot, facts in [("MAIN", []), ("PT01", ["purchase.mobility", "display.size"]), ("PT02", ["display.size", "work.scene"])]]
    observations.write_text(json.dumps({"images": items}))
    findings = audit_images(RECORD, replace(clean_evidence(), image_urls=[str(path)] * 3), observations_path=observations)
    overlaps = [f for f in findings if f.rule_id == "IMG-SEMANTIC-002"]
    assert len(overlaps) == 1
    assert overlaps[0].observed == "50.0%: display.size"


def test_cli_caches_erp_per_product_and_skipped_images_stay_review(monkeypatch, tmp_path):
    from listing_auditor import cli
    records = [RECORD, replace(RECORD, seller="JTD")]
    calls = []
    results = []
    monkeypatch.setattr("sys.argv", ["listing-auditor", "audit", "--input", "unused.csv", "--erp-export", "unused.json", "--skip-image-audit", "--out", str(tmp_path)])
    monkeypatch.setattr(cli, "load_records", lambda **kwargs: records)
    def load(sku, **kwargs):
        calls.append(sku)
        return ERPRecord(sku, available=True)
    monkeypatch.setattr(cli, "load_erp_record", load)
    monkeypatch.setattr(cli, "collect_evidence", lambda *args: clean_evidence())
    def write(items, root):
        results.extend(items)
        return [root / "test.xlsx"]
    monkeypatch.setattr(cli, "write_product_packages", write)
    cli.main()
    assert calls == [RECORD.sku]
    assert all(any(f.rule_id == "IMG-SKIPPED-001" for f in result.findings) for result in results)


def test_unresolved_erp_conflict_does_not_become_confident_copy():
    from listing_auditor.models import Finding
    conflicts = [Finding("memory", "16GB", "32GB", "REVIEW", "ERP conflict", corrected_value="MANUAL REVIEW")]
    conflicts.extend(check_compliance(RECORD, replace(clean_evidence(), title="Wrong")))
    erp = ERPRecord(RECORD.sku, {"ram": "32GB RAM"}, available=True)
    drafts = style_correction_findings(RECORD, clean_evidence(), erp, conflicts)
    assert "[[VERIFY CONFLICTING MEMORY]]" in next(f.corrected_value for f in drafts if f.field == "copy_title")


def test_export_preserves_bullets_and_versioned_metadata(tmp_path):
    from openpyxl import load_workbook
    from listing_auditor.report import write_product_packages
    result = compare(RECORD, clean_evidence())
    path = write_product_packages([result], tmp_path)[0]
    book = load_workbook(path, read_only=True)
    rows = list(book["Source Evidence"].iter_rows(min_row=2, values_only=True))
    assert [row[4] for row in rows if row[3].startswith("Bullet ")] == clean_evidence().bullets
    output = json.loads((path.parent / "report.json").read_text(encoding="utf-8"))
    assert output[0]["audit_metadata"]["auditor_version"] == "0.8.0"
    assert output[0]["audit_metadata"]["publication_eligibility"] == "NOT_ASSESSED"


def test_malformed_image_observations_require_review(tmp_path):
    image = tmp_path / "main.png"
    Image.new("RGB", (1000, 1000), "white").save(image)
    observations = tmp_path / "observations.json"
    observations.write_text('{"images": [null]}')
    findings = audit_images(RECORD, replace(clean_evidence(), image_urls=[str(image)]), observations_path=observations)
    assert any(f.rule_id == "IMG-EVIDENCE-002" for f in findings)
