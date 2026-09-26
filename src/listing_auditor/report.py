from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import AuditResult


HEADERS = [
    "严重程度", "问题类型", "Rule ID / Reference", "产品名称", "内部编号", "Seller", "ASIN",
    "影响位置 / 字段", "错误或歧义内容", "已核实 / 预期内容",
    "建议修改内容", "问题说明", "Amazon Reference", "Input / ERP Reference",
    "证据摘录", "人工审核状态", "人工最终修改 / 备注", "审核人", "审核日期",
]


def _issue_type(field: str, severity: str) -> str:
    if field.startswith("compliance"):
        return "合规风险"
    if field.startswith("image_"):
        return "图片问题"
    if field.startswith("copy_"):
        return "文案修改建议"
    if field in {"brand", "model"}:
        return "上下文不对应"
    if field in {"evidence", "erp_evidence"} or severity == "REVIEW":
        return "证据不足 / 有歧义"
    return "具体规格错误"


def _rows(results: list[AuditResult]) -> list[list[str]]:
    rows: list[list[str]] = []
    for result in results:
        for finding in result.findings:
            rows.append([
                finding.severity,
                _issue_type(finding.field, finding.severity),
                " — ".join(value for value in [finding.rule_id, finding.reference] if value),
                result.listing.product_name or result.evidence.title or f"Amazon ASIN {result.listing.asin}",
                result.listing.sku,
                result.listing.seller,
                result.listing.asin,
                finding.field,
                finding.observed,
                finding.erp_value or finding.expected,
                finding.corrected_value,
                finding.reason,
                result.listing.url,
                f"Input: {finding.expected}; ERP: {finding.erp_value or 'not confirmed'}",
                f"{finding.evidence_source}: {finding.observed}",
                "待审核",
                "",
                "",
                "",
            ])
    return rows


def _write_excel(results: list[AuditResult], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "问题清单"
    issue_count = sum(len(r.findings) for r in results)
    sheet["A2"] = "Amazon Listing 审核问题清单"
    sheet.merge_cells("A2:S2")
    sheet["A3"] = "仅包含发现错误、上下文不对应、合规风险或证据不足的项目。已确认无问题的 listing 不会出现在此表。"
    sheet.merge_cells("A3:S3")
    sheet["A5"], sheet["B5"] = "问题产品", sum(bool(r.findings) for r in results)
    sheet["D5"], sheet["E5"] = "问题 / 歧义", issue_count
    sheet["G5"], sheet["H5"] = "Critical", sum(f.severity == "CRITICAL" for r in results for f in r.findings)
    sheet["J5"], sheet["K5"] = "待人工审核", issue_count
    for index, header in enumerate(HEADERS, 1):
        sheet.cell(row=8, column=index, value=header)
    for row in _rows(results):
        sheet.append(row)

    navy = "17365D"
    sheet.sheet_view.showGridLines = False
    sheet["A2"].font = Font(name="Arial", size=16, bold=True, color="1F2937")
    sheet["A3"].font = Font(name="Arial", size=10, italic=True, color="64748B")
    for start in ("A5:B5", "D5:E5", "G5:H5", "J5:K5"):
        for row in sheet[start]:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="EAF2F8")
                cell.font = Font(name="Arial", size=10, bold=True, color=navy)
    for cell in sheet[8]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    for row in sheet.iter_rows(min_row=9):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "E9"
    sheet.auto_filter.ref = f"A8:S{max(sheet.max_row, 8)}"
    if sheet.max_row > 8:
        table = Table(displayName="AuditIssues", ref=f"A8:S{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
        sheet.add_table(table)
    widths = [11, 27, 34, 42, 15, 11, 14, 24, 34, 32, 34, 44, 34, 38, 44, 18, 38, 18, 16]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    yellow = PatternFill("solid", fgColor="FFF2CC")
    for row in sheet.iter_rows(min_row=9, min_col=16, max_col=19):
        for cell in row:
            cell.fill = yellow
    validation = DataValidation(type="list", formula1='"待审核,确认错误,已修改,False Positive"', allow_blank=False)
    sheet.add_data_validation(validation)
    validation.add(f"P9:P{max(sheet.max_row, 9)}")
    for formula, color in [("$A9=\"CRITICAL\"", "E06666"), ("$A9=\"HIGH\"", "FCE5CD"), ("$A9=\"REVIEW\"", "FFF2CC")]:
        sheet.conditional_formatting.add(f"A9:S{max(sheet.max_row, 9)}", FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=color)))

    evidence_sheet = workbook.create_sheet("Source Evidence")
    evidence_sheet.append(["Internal ID", "ASIN", "Source", "Field / Section", "Value", "URL"])
    for result in (result for result in results if result.findings):
        evidence_sheet.append([result.listing.sku, result.listing.asin, result.evidence.source, "Title", result.evidence.title, result.evidence.url])
        evidence_sheet.append([result.listing.sku, result.listing.asin, result.evidence.source, "Description", result.evidence.description, result.evidence.url])
        for key, value in result.evidence.details.items():
            evidence_sheet.append([result.listing.sku, result.listing.asin, result.evidence.source, key, value, result.evidence.url])
        for index, value in enumerate(result.evidence.image_urls[:9]):
            slot = "MAIN" if index == 0 else f"PT{index:02d}"
            evidence_sheet.append([result.listing.sku, result.listing.asin, result.evidence.source, f"Image {slot}", value, value])
        if result.erp:
            for key, value in result.erp.fields.items():
                evidence_sheet.append([result.listing.sku, result.listing.asin, result.erp.source, key, value, ""])
    for cell in evidence_sheet[1]:
        cell.font = Font(name="Arial", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=navy)
    evidence_sheet.freeze_panes = "A2"
    evidence_sheet.auto_filter.ref = evidence_sheet.dimensions
    for column, width in zip("ABCDEF", [16, 14, 24, 24, 80, 42]):
        evidence_sheet.column_dimensions[column].width = width
    for row in evidence_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    workbook.save(path)


def write_reports(results: list[AuditResult], out: Path) -> Path:
    out.mkdir(parents=True, exist_ok=True)
    (out / "report.json").write_text(json.dumps([r.to_dict() for r in results], indent=2, ensure_ascii=False), encoding="utf-8")
    with (out / "report.csv").open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        writer.writerows(_rows(results))

    lines = ["# Amazon Listing Audit Report", "", f"Listings checked: {len(results)}", ""]
    for result in results:
        lines.extend([f"## {result.status}: {result.listing.asin} ({result.listing.seller})", "", f"- Expected: {result.listing.product_name}", f"- Internal ID: {result.listing.sku}", f"- URL: {result.listing.url}"])
        if not result.findings:
            lines.append("- No mismatch found in available evidence.")
        for finding in result.findings:
            lines.append(f"- **{finding.severity} / {finding.field}** — input `{finding.expected}`, Amazon `{finding.observed}`, ERP `{finding.erp_value}`. {finding.reason}")
        lines.append("")
    (out / "report.md").write_text("\n".join(lines), encoding="utf-8")
    excel_path = out / "listing-audit-review.xlsx"
    _write_excel(results, excel_path)
    return excel_path
