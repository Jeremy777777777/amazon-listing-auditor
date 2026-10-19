from importlib.resources import files
import hashlib
import json
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from listing_auditor.models import AuditResult, Evidence, Finding, ListingRecord
from listing_auditor.template_report import TEMPLATE_SHEETS, write_template_workbook


def result(asin="B012345678", title="Original title", findings=None):
    record = ListingRecord("Dell QC24251 / Intel i5-14500T", "VA-215", "MegaPC", f"https://www.amazon.com/dp/{asin}", asin, 1)
    evidence = Evidence(record.url, title=title, bullets=["Bullet one", "Bullet two"], description="Original description", details={"Brand": "MegaPC", "Screen Size": "23.8 inches", "Total Number of HDMI Ports": "1"}, source="test")
    return AuditResult(record, "FAIL", findings or [Finding("copy_title", "Correct title", title, "HIGH", "Bad title", corrected_value="Recommended title", rule_id="STYLE-TITLE-001")], evidence)


def test_source_template_is_preserved_and_bad_value_is_red(tmp_path: Path):
    path = tmp_path / "audit.xlsx"
    write_template_workbook([result()], path)
    wb = load_workbook(path)
    source = load_workbook(BytesIO(files("listing_auditor").joinpath("assets/listing-workbook-template.xlsx").read_bytes()))
    for name in TEMPLATE_SHEETS:
        assert wb[name].max_column == source[name].max_column
        assert [c.value for c in wb[name][1]] == [c.value for c in source[name][1]]
        assert [c.value for c in wb[name][2]] == [c.value for c in source[name][2]]
    product = wb["Product Details"]
    assert product["C3"].value == "Original title"
    assert product["C3"].font.color.rgb == "009C0006"
    assert product["C6"].value == "Recommended title"
    assert "STYLE-TITLE-001" in product["C7"].value
    assert product["E3"].value == "MegaPC"
    assert product["E4"].value == "CAPTURED_NOT_FLAGGED"
    assert product["F3"].value is None
    assert product["F4"].value == "NOT_CAPTURED"
    assert wb["Offer"]["B3"].value is None  # internal ID is not a verified merchant SKU
    assert wb["审核修改"]["A2"].hyperlink.target == "#'Product Details'!C3"
    assert wb["审核修改"]["C2"].value == "Original title"
    assert wb["审核修改"]["D2"].value == "Recommended title"
    assert any(str(v.sqref) == "AF3" for v in wb["Safety&Compliance"].data_validations.dataValidation)
    assert "ISSUE_REVIEW" in product.data_validations.dataValidation[0].formula1
    assert "D" not in product.column_dimensions  # preserve source grouped widths
    provenance = json.loads(files("listing_auditor").joinpath("assets/template-provenance.json").read_text(encoding="utf-8"))
    assert hashlib.sha256(files("listing_auditor").joinpath("assets/listing-workbook-template.xlsx").read_bytes()).hexdigest() == provenance["sha256"]


def test_alias_mapping_image_findings_and_unknown_rules_remain_visible(tmp_path: Path):
    findings = [Finding("Total HDMI Ports", "0", "1", "HIGH", "Wrong port count", corrected_value="0"), Finding("evidence", "Readable page", "Timeout", "REVIEW", "Missing page"), Finding("image_PT01_cpu", "i5-14500T", "i5-1334U", "HIGH", "Wrong CPU in image", corrected_value="Use i5-14500T")]
    path = tmp_path / "audit.xlsx"
    write_template_workbook([result(findings=findings)], path)
    wb = load_workbook(path)
    p = wb["Product Details"]
    column = next(c.column for c in p[1] if c.value == "Total HDMI Ports")
    assert p.cell(3, column).value == "1"
    assert p.cell(6, column).value == "0"
    assert any(c.value == "Timeout" for row in wb["审核修改"] for c in row)
    assert wb["图片审核"]["C2"].value == "i5-1334U"
    assert p["C4"].value == "CAPTURED_NOT_FLAGGED"  # image error does not flag unrelated title


def test_multiple_listings_do_not_overwrite_each_other_and_formula_text_is_literal(tmp_path: Path):
    path = tmp_path / "audit.xlsx"
    write_template_workbook([result(title="=HYPERLINK(\"https://example.invalid\",\"title\")"), result(asin="B087654321", title="Second ASIN")], path)
    wb = load_workbook(path)
    assert wb["Product Details"]["C3"].data_type == "s"
    assert wb["Product Details 2"]["C3"].value == "Second ASIN"
    assert wb["Offer 2"]["B3"].value is None
    assert wb["Product Details 2"]["F4"].value == "NOT_CAPTURED"


def test_long_copy_is_not_truncated_in_reading_view(tmp_path: Path):
    long = "Full paragraph " * 300
    finding = Finding("copy_product_description", "Correct text", long, "HIGH", "Needs correction", corrected_value=long + "new")
    r = result(findings=[finding])
    object.__setattr__(r.evidence, "description", long)
    path = tmp_path / "audit.xlsx"
    write_template_workbook([r], path)
    wb = load_workbook(path)
    assert wb["Product Details"]["L3"].value == long
    assert wb["Product Details"]["L6"].value == long + "new"
    review = wb["审核修改"]
    assert "".join(str(review.cell(row, 3).value or "") for row in range(2, review.max_row + 1)) == long
    assert "".join(str(review.cell(row, 4).value or "") for row in range(2, review.max_row + 1)) == long + "new"
    assert review.data_validations.dataValidation[0].sqref == f"F2:F{review.max_row}"
