"""Cell-level review on the pinned generation workbook, without changing source values."""
from __future__ import annotations

import math
import re
from copy import copy
from importlib.resources import files
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .models import AuditResult, Finding


TEMPLATE_SHEETS = ("Product Details", "Offer", "Safety&Compliance")
RED = "9C0006"
RED_FILL = "FCE4D6"
AMBER = "FFF2CC"
GREEN = "E2EFDA"
NAVY = "17365D"


def _norm(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.casefold())


# Amazon display labels differ from the generation template. No fuzzy product matching.
ALIASES = {
    "Item Name": ("Title",), "Brand Name": ("Brand",),
    "CPU Model": ("Processor Series", "CPU Model"), "Cpu Model": ("Processor Series", "CPU Model"),
    "CPU Model Manufacturer": ("Processor Brand",), "CPU Model Number": ("CPU Model Number",),
    "Graphics Description": ("Graphics Card Description", "Graphics Description"),
    "Graphics Card Ram": ("Graphics Card Ram",),
    "RAM Memory Installed": ("Memory / RAM Memory Installed", "RAM Memory Installed"),
    "RAM Memory Maximum Size": ("Ram Memory Maximum Size",),
    "Ram Memory Maximum Size": ("Ram Memory Maximum Size",),
    "Total HDMI Ports": ("Total Number of HDMI Ports",),
    "HDMI Port Quantity": ("Total Number of HDMI Ports",),
    "Number of HDMI ports": ("Total Number of HDMI Ports",),
    "Total USB Ports": ("Total Usb Ports",),
    "Specific Uses for Product": ("Specific Uses For Product",),
    "Human Interface Input": ("Human-Interface Input",),
    "Included Components": ("Built-In Media", "Included Components"),
    "Form Factor": ("Personal Computer Design Type",),
    "Hard Disk Size": ("Hard Disk Size",),
    "Customizations": ("Customizations",),
}

FIELD_MAP = {
    "brand": ("Brand Name", "Manufacturer"), "model": ("Model Number", "Model Name"),
    "cpu": ("CPU Model", "Cpu Model", "CPU Model Number"),
    "resolution": ("Display Maximum Resolution", "Display Resolution Maximum", "Native Resolution", "Resolution"),
    "display_size": ("Screen Size",), "refresh_rate": ("Display Refresh Rate in Hertz",),
    "memory": ("RAM Memory Installed",), "storage": ("Hard Disk Size", "Memory Storage Capacity"),
    "os": ("Operating System",), "touchscreen": ("Touch Screen Type",),
    "bluetooth": ("Bluetooth support?",), "wifi": ("Connectivity Technology",),
    "fingerprint": ("Biometric Security Feature",), "backlit": ("Keyboard Description",),
}


def _targets(finding: Finding) -> set[str]:
    """Map findings to affected template fields; unmatched findings remain on the review page."""
    field = finding.field
    if field.startswith("image_"):
        return set()
    if field in FIELD_MAP:
        return set(FIELD_MAP[field])
    if field == "copy_title" or field.startswith("compliance_title"):
        return {"Item Name"}
    if field == "copy_bullet_points":
        return {"Bullet Point"}
    if field == "copy_product_description":
        return {"Product Description"}
    if "identity" in field:
        if "description" in field: return {"Product Description"}
        if "bullet" in field: return {"Bullet Point"}
        return {"Model Name", "Model Number"}
    if "warranty" in field.casefold() or "WARRANTY" in finding.rule_id:
        return {"Bullet Point", "Product Description", "Warranty Description"}
    if field.startswith("compliance"):
        return {"Bullet Point", "Product Description"}
    if finding.rule_id == "MANUAL-CONFIG-001":
        return {"Item Name", "Bullet Point", "Product Description", "RAM Memory Installed"}
    if finding.rule_id == "MANUAL-CONFIG-002":
        return {"Item Name", "Bullet Point", "Product Description", "Hard Disk Size", "Memory Storage Capacity"}
    if finding.rule_id == "MANUAL-SPEC-001": return {"Graphics Card Ram", "Graphics Description"}
    if finding.rule_id == "MANUAL-SPEC-002": return {"Cpu Model", "CPU Model", "Processor Speed"}
    if finding.rule_id == "MANUAL-PORTS-001": return {"Total HDMI Ports", "HDMI Port Quantity", "Number of HDMI ports"}
    # Exact display label / template label, never an arbitrary substring match.
    return {field.split(" — ", 1)[-1]}


def _literal(cell, value) -> None:
    """Untrusted listing copy is always text, even if it starts with '='."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _append_literals(sheet, values):
    row = sheet.max_row + 1
    for column, value in enumerate(values, 1):
        _literal(sheet.cell(row, column), value)


def _chunks(value, limit=1000):
    """Split only the reading view; the template and evidence retain exact complete strings."""
    text = str(value or "")
    if not text: return [""]
    parts = []
    while len(text) > limit:
        cut = text.rfind("\n\n", 0, limit)
        if cut >= limit // 2:
            cut += 2
        else:
            cut = text.rfind(" ", 0, limit)
            cut = cut + 1 if cut >= limit // 2 else limit
        parts.append(text[:cut])
        text = text[cut:]
    parts.append(text)
    return parts


def _source_value(result: AuditResult, sheet: str, field: str):
    evidence = result.evidence
    if sheet == "Product Details" and field == "Item Name": return evidence.title or None
    if sheet == "Product Details" and field == "Bullet Point": return "\n\n".join(evidence.bullets) or None
    if sheet == "Product Details" and field == "Product Description": return evidence.description or None
    # Internal ID is not necessarily an Amazon merchant SKU; Offer values require evidence.
    candidates = ALIASES.get(field, (field,))
    for name in candidates:
        for key, value in evidence.details.items():
            if _norm(key) == _norm(name): return value
    return None


def _recommendation(findings: list[Finding], field: str) -> str:
    copy_field = {"Item Name": "copy_title", "Bullet Point": "copy_bullet_points", "Product Description": "copy_product_description"}.get(field)
    # A full proposed copy block takes precedence over many short edit instructions.
    draft = next((f.corrected_value for f in findings if f.field == copy_field and f.corrected_value), None)
    if draft: return draft
    return "\n\n".join(dict.fromkeys(f.corrected_value or f.expected or "待核实：请补充依据后修改" for f in findings))


def _format_table(sheet, widths, header_row=1):
    sheet.sheet_view.showGridLines = False
    for c, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, c).column_letter].width = width
    for row in sheet.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="1F2937")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for cell in sheet[header_row]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", size=10, color="FFFFFF", bold=True)
    sheet.row_dimensions[header_row].height = 30
    sheet.freeze_panes = f"C{header_row + 1}"
    for r in range(header_row + 1, sheet.max_row + 1):
        lines = 1
        for c, width in enumerate(widths, 1):
            text = str(sheet.cell(r, c).value or "")
            count = sum(max(1, math.ceil(sum(1.8 if ord(ch) > 255 else 1 for ch in line) / max(8, width - 3))) for line in text.split("\n"))
            lines = max(lines, count)
        sheet.row_dimensions[r].height = min(400, max(32, lines * 14 + 10))
    if sheet.max_row > header_row:
        sheet.auto_filter.ref = f"A{header_row}:{sheet.cell(sheet.max_row, len(widths)).coordinate}"


def _template_heights(sheet):
    for row in (3, 5, 6, 7):
        lines = 1
        for column in range(1, sheet.max_column + 1):
            cell = sheet.cell(row, column)
            # Reading a missing ColumnDimension via [] creates a new 13-char override;
            # respect the source template's grouped column widths instead.
            dimension = sheet.column_dimensions.get(cell.column_letter)
            if dimension is None:
                dimension = next((item for item in sheet.column_dimensions.values() if item.min and item.max and item.min <= column <= item.max), None)
            width = dimension.width if dimension is not None else 24
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            text = str(cell.value or "")
            lines = max(lines, sum(max(1, math.ceil(sum(1.8 if ord(ch) > 255 else 1 for ch in part) / max(8, width - 3))) for part in text.split("\n")))
        sheet.row_dimensions[row].height = min(400, max(32, lines * 14 + 10))


def _write_review_rows(sheet, rows):
    headers = ["位置 / 模板单元格", "字段 / 规则", "Amazon 原内容（红色为问题）", "推荐修改值 / 内容", "问题说明 / 依据", "人工审核状态", "最终修改 / 备注"]
    sheet.append(headers)
    expanded = []
    for row in rows:
        current, recommended = _chunks(row[2]), _chunks(row[3])
        for i in range(max(len(current), len(recommended))):
            expanded.append([row[0], row[1] + (f"\n内容续页 {i + 1}" if i else ""), current[i] if i < len(current) else "", recommended[i] if i < len(recommended) else "", row[4] if i == 0 else "同一字段的连续文本；完整内容在模板原单元格。", row[5], row[6]])
    for row in expanded:
        sheet.append([None] * len(headers))
        r = sheet.max_row
        for c, value in enumerate(row, 1): _literal(sheet.cell(r, c), value)
        sheet.cell(r, 3).fill = PatternFill("solid", fgColor=RED_FILL)
        sheet.cell(r, 4).fill = PatternFill("solid", fgColor=GREEN)
        for c in (6, 7): sheet.cell(r, c).fill = PatternFill("solid", fgColor=AMBER)
    _format_table(sheet, [27, 35, 65, 75, 60, 18, 45])
    for r in range(2, sheet.max_row + 1):
        sheet.cell(r, 3).font = Font(name="Arial", size=10, color=RED)
    if rows:
        dropdown = DataValidation(type="list", formula1='"待审核,确认错误,已修改,False Positive"')
        sheet.add_data_validation(dropdown)
        dropdown.add(f"F2:F{sheet.max_row}")


def write_template_workbook(results: list[AuditResult], path: Path) -> None:
    results = [result for result in results if result.findings] or results
    template_bytes = files("listing_auditor").joinpath("assets/listing-workbook-template.xlsx").read_bytes()
    workbook = load_workbook(BytesIO(template_bytes))
    review = workbook.create_sheet("审核修改", 0)
    image_review = workbook.create_sheet("图片审核")
    evidence_sheet = workbook.create_sheet("Source Evidence")
    evidence_sheet.append(["Internal ID", "ASIN / Seller", "Source", "Field / Section", "Value", "URL"])
    review_rows, image_rows = [], []
    image_links = {}
    for index, result in enumerate(results):
        # Preserve one complete template per listing, including several ASINs for one internal ID.
        if index == 0:
            sheets = [workbook[name] for name in TEMPLATE_SHEETS]
        else:
            clean = load_workbook(BytesIO(template_bytes))
            # copy_worksheet is deliberately within the output workbook; reset values from a clean template.
            sheets = []
            for name in TEMPLATE_SHEETS:
                target = workbook.copy_worksheet(workbook[name])
                for row in clean[name]:
                    for cell in row:
                        target.cell(cell.row, cell.column).value = cell.value
                        target.cell(cell.row, cell.column)._style = copy(cell._style)
                target.title = f"{name[:24]} {index + 1}"
                sheets.append(target)
        assigned: set[int] = set()
        label = f"{result.listing.asin} / {result.listing.seller}"
        for sheet in sheets:
            original_name = TEMPLATE_SHEETS[sheets.index(sheet)]
            sheet.sheet_view.showGridLines = False
            sheet.freeze_panes = "B3"
            sheet.cell(6, 1).value = "推荐修改"
            sheet.cell(7, 1).value = "问题说明 / Rule ID"
            # Source template status dropdowns use production states; replace them with audit states.
            sheet.data_validations.dataValidation = [validation for validation in sheet.data_validations.dataValidation if not any(area.min_row <= 4 <= area.max_row for area in validation.sqref.ranges)]
            status_validation = DataValidation(type="list", formula1='"NOT_CAPTURED,CAPTURED_NOT_FLAGGED,ISSUE_REVIEW"')
            sheet.add_data_validation(status_validation)
            status_validation.add(f"B4:{sheet.cell(4, sheet.max_column).coordinate}")
            for column in range(2, sheet.max_column + 1):
                field = str(sheet.cell(1, column).value or "")
                current = _source_value(result, original_name, field)
                _literal(sheet.cell(3, column), current)
                sheet.cell(4, column).value = "CAPTURED_NOT_FLAGGED" if current is not None else "NOT_CAPTURED"
                _literal(sheet.cell(5, column), f"{result.evidence.source}\n{result.listing.url}" if current is not None else "未取得该字段的页面证据；不代表通过")
                sheet.cell(6, column).value = None
                sheet.cell(7, column).value = None
                matches = [(i, f) for i, f in enumerate(result.findings) if field in _targets(f)]
                if not matches: continue
                assigned.update(i for i, _ in matches)
                field_findings = [f for _, f in matches]
                recommendation = _recommendation(field_findings, field)
                reason = "\n\n".join(dict.fromkeys(f"{f.severity} / {f.rule_id or f.field}: {f.reason}" for f in field_findings))
                source_cell = sheet.cell(3, column)
                source_cell.fill = PatternFill("solid", fgColor=RED_FILL)
                source_cell.font = Font(name="Arial", size=10, color=RED)
                sheet.cell(4, column).value = "ISSUE_REVIEW"
                _literal(sheet.cell(6, column), recommendation)
                sheet.cell(6, column).fill = PatternFill("solid", fgColor=GREEN)
                _literal(sheet.cell(7, column), reason)
                for r in (3, 5, 6, 7): sheet.cell(r, column).alignment = Alignment(wrap_text=True, vertical="top")
                if field in {"Item Name", "Bullet Point", "Product Description"}: sheet.column_dimensions[source_cell.column_letter].width = 80
                address = f"'{sheet.title}'!{source_cell.coordinate}"
                summary = " / ".join(dict.fromkeys(f.rule_id or f.field for f in field_findings))
                summary += "\n" + field_findings[0].reason + "\n完整问题说明见模板第 7 行。"
                review_rows.append([address, f"{label}\n{field}", current if current is not None else "未抓取；问题依据见右侧", recommendation, summary, "待审核", ""])
            _template_heights(sheet)
        for i, finding in enumerate(result.findings):
            if i in assigned: continue
            row = [finding.field, f"{label}\n{finding.rule_id or finding.field}", finding.observed, finding.corrected_value or finding.expected, f"{finding.severity}: {finding.reason}\n{finding.reference}", "待审核", ""]
            (image_rows if finding.field.startswith("image_") else review_rows).append(row)
            if finding.field.startswith("image_") and finding.evidence_source.startswith(("https://", "http://")):
                image_links[(row[0], row[1])] = finding.evidence_source
        # Evidence is exhaustive, field-separated, and includes Bullet Points and collection failures.
        entries = [("Audit status", result.status), ("Input product", result.listing.product_name), ("Internal ID", result.listing.sku), ("Title", result.evidence.title), ("Captured at", result.evidence.captured_at), ("Availability", str(result.evidence.available)), ("Collection error", result.evidence.error)]
        entries += [(f"Bullet {i}", value) for i, value in enumerate(result.evidence.bullets, 1)]
        entries += [(f"Description {i}", value) for i, value in enumerate(result.evidence.description.split("\n\n"), 1)]
        entries += list(result.evidence.details.items())
        entries += [("MAIN" if i == 0 else f"PT{i:02d}", url) for i, url in enumerate(result.evidence.image_urls)]
        for field, value in entries:
            source = "User input" if field in {"Input product", "Internal ID"} else ("Audit" if field == "Audit status" else result.evidence.source)
            _append_literals(evidence_sheet, [result.listing.sku, label, source, field, None, result.listing.url])
            _literal(evidence_sheet.cell(evidence_sheet.max_row, 5), value)
        if result.erp:
            for field, value in [("ERP available", str(result.erp.available)), ("ERP error", result.erp.error), *result.erp.fields.items()]:
                _append_literals(evidence_sheet, [result.listing.sku, label, result.erp.source, field, None, ""])
                _literal(evidence_sheet.cell(evidence_sheet.max_row, 5), value)
    _write_review_rows(review, review_rows)
    _write_review_rows(image_review, image_rows)
    for r in range(2, image_review.max_row + 1):
        key = (image_review.cell(r, 1).value, str(image_review.cell(r, 2).value or "").split("\n内容续页", 1)[0])
        if key in image_links: image_review.cell(r, 3).hyperlink = image_links[key]
    _format_table(evidence_sheet, [16, 35, 30, 40, 95, 50])
    # Navigation is an actual workbook link; source and image links retain their public URLs.
    for r in range(2, review.max_row + 1):
        location = str(review.cell(r, 1).value or "")
        if location.startswith("'"): review.cell(r, 1).hyperlink = "#" + location
    for sheet in (review, image_review):
        if sheet.max_row == 1:
            sheet.append(["没有记录到问题；仅表示可用证据未触发规则，不代表完整审核通过。"])
    for r in range(2, evidence_sheet.max_row + 1):
        url = evidence_sheet.cell(r, 6)
        if str(url.value or "").startswith(("https://", "http://")): url.hyperlink = str(url.value)
    workbook.active = 0
    workbook.save(path)
