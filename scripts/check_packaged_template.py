"""Check the distributable, not the editable checkout, for its pinned XLSX resource."""
from __future__ import annotations

import hashlib
import json
import sys
import zipfile
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    wheel_path = Path(sys.argv[1])
    with zipfile.ZipFile(wheel_path) as wheel:
        template = wheel.read("listing_auditor/assets/listing-workbook-template.xlsx")
        provenance = json.loads(wheel.read("listing_auditor/assets/template-provenance.json"))
    assert hashlib.sha256(template).hexdigest() == provenance["sha256"]
    workbook = load_workbook(BytesIO(template))
    assert workbook.sheetnames == ["Product Details", "Offer", "Safety&Compliance"]
    assert workbook["Product Details"]["C1"].value == "Item Name"
    print(f"Packaged template verified: {wheel_path.name}")


if __name__ == "__main__":
    main()
